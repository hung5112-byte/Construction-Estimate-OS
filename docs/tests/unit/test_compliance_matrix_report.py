from openpyxl import load_workbook

from core.tools.compliance_tracker import ComplianceTracker
from core.tools.compliance_matrix_report import build_compliance_matrix, score_item

AS_OF = "2026-07-01"
ITEMS = [
    {"product": "AMG-100", "cert": "RoHS", "expiry": "2027-03-01", "required": "yes"},   # Valid
    {"product": "AMG-100", "cert": "FCC", "expiry": "2026-06-01", "required": "yes"},    # Expired -> block
    {"product": "AMG-100", "cert": "CE", "status": "missing", "required": "yes"},        # Missing -> block
    {"product": "CY-80", "cert": "UL", "expiry": "2026-08-15", "required": "yes"},       # Expiring soon
]


def test_score_item_matches_tool():
    tool = ComplianceTracker().run(f"cert FCC expiry=2026-06-01 required=yes as_of={AS_OF}").data
    rep = score_item({**ITEMS[1], "as_of": AS_OF})
    assert rep["status"] == tool["status"]
    assert rep["shippable"] == tool["shippable"]


def test_report_has_both_sheets(tmp_path):
    out = build_compliance_matrix(ITEMS, tmp_path / "cm.xlsx", as_of=AS_OF)
    assert out.exists()
    assert load_workbook(out).sheetnames == ["Summary", "Items"]


def test_items_sorted_worst_first(tmp_path):
    out = build_compliance_matrix(ITEMS, tmp_path / "cm.xlsx", as_of=AS_OF)
    ws = load_workbook(out)["Items"]
    tiers = [ws.cell(r, 5).value for r in range(5, 5 + len(ITEMS))]
    # Critical (Expired/Missing) must come before Low (Valid)
    assert tiers[0] == "Critical"
    assert tiers[-1] == "Low"


def test_summary_counts_blocking(tmp_path):
    out = build_compliance_matrix(ITEMS, tmp_path / "cm.xlsx", as_of=AS_OF)
    ws = load_workbook(out)["Summary"]
    flat = {row[0]: row[1] for row in ws.iter_rows(min_row=5, max_col=2, values_only=True) if row[0]}
    assert flat["Total items"] == 4
    assert flat["Shipment-blocking"] == 2   # FCC expired + CE missing
    assert flat["Expiring soon"] == 1


def test_empty_raises(tmp_path):
    try:
        build_compliance_matrix([], tmp_path / "x.xlsx")
        assert False, "expected ValueError"
    except ValueError:
        pass
