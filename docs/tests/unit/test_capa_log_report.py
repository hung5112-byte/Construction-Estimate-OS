from openpyxl import load_workbook

from core.tools.capa_8d import Capa8D
from core.tools.capa_log_report import build_capa_log_report, score_capa

CAPAS = [
    {"capa": "Q-1", "step": 2, "days_open": 3, "severity": "critical"},          # At Risk, escalate
    {"capa": "Q-2", "step": 5, "days_open": 80, "sla_days": 60, "severity": "high", "containment": "yes"},  # Overdue
    {"capa": "Q-3", "step": 3, "days_open": 5, "severity": "medium", "containment": "yes"},  # On Track
    {"capa": "Q-4", "step": 8, "days_open": 40, "severity": "medium"},           # Closed
]


def test_score_capa_matches_tool():
    tool = Capa8D().run("capa Q-2 step=5 days_open=80 sla_days=60 severity=high containment=yes").data
    rep = score_capa(CAPAS[1])
    assert rep["status"] == tool["status"]
    assert rep["priority"] == tool["priority"]


def test_report_has_both_sheets(tmp_path):
    out = build_capa_log_report(CAPAS, tmp_path / "capa.xlsx")
    assert out.exists()
    assert load_workbook(out).sheetnames == ["Summary", "Log"]


def test_log_is_ordered_overdue_first_closed_last(tmp_path):
    out = build_capa_log_report(CAPAS, tmp_path / "capa.xlsx")
    ws = load_workbook(out)["Log"]
    statuses = [ws.cell(r, 2).value for r in range(5, 5 + len(CAPAS))]
    assert statuses[0] == "Overdue"
    assert statuses[-1] == "Closed"


def test_summary_counts(tmp_path):
    out = build_capa_log_report(CAPAS, tmp_path / "capa.xlsx")
    ws = load_workbook(out)["Summary"]
    flat = {row[0]: row[1] for row in ws.iter_rows(min_row=5, max_col=2, values_only=True) if row[0]}
    assert flat["Total CAPAs"] == 4
    assert flat["Overdue"] == 1
    assert flat["Closed"] == 1
    assert flat["Needs escalation"] >= 2  # Q-1 (containment gap) + Q-2 (overdue)


def test_empty_raises(tmp_path):
    try:
        build_capa_log_report([], tmp_path / "x.xlsx")
        assert False, "expected ValueError"
    except ValueError:
        pass
