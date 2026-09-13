from openpyxl import load_workbook

from core.tools.landed_cost import LandedCost
from core.tools.landed_cost_workbook import (
    build_landed_cost_workbook,
    compute_scenario,
)

CHINA = {
    "name": "China", "ex_works": 42.50, "units": 5000, "duty": 2.6,
    "tariff": 25, "freight": 9000, "hmf": 0.125, "broker": 250,
    "transit_days": 32, "cost_of_capital": 9,
}
MEXICO = {
    "name": "Mexico (USMCA)", "ex_works": 46.00, "units": 5000, "duty": 0,
    "tariff": 0, "freight": 2500, "broker": 200, "transit_days": 5,
    "cost_of_capital": 9,
}


def test_compute_scenario_matches_tool():
    # the workbook's numbers must equal the landed_cost tool exactly
    tool = LandedCost().run(
        "landed 42.50 units=5000 duty=2.6 tariff=25 freight=9000 hmf=0.125 "
        "broker=250 transit_days=32 cost_of_capital=9"
    ).data
    wb = compute_scenario(CHINA)
    assert wb["landed_unit"] == tool["landed_unit"]
    assert wb["tariff_unit"] == tool["tariff_unit"]


def test_workbook_has_all_expected_sheets(tmp_path):
    out = build_landed_cost_workbook([CHINA, MEXICO], tmp_path / "lc.xlsx")
    assert out.exists()
    wb = load_workbook(out)
    assert "Inputs" in wb.sheetnames
    assert "Comparison" in wb.sheetnames
    assert "Sensitivity" in wb.sheetnames
    # one sheet per scenario (names sanitized/truncated to <=28 chars)
    assert "China" in wb.sheetnames
    assert any(s.startswith("Mexico") for s in wb.sheetnames)


def test_scenario_sheet_landed_value_matches_engine(tmp_path):
    out = build_landed_cost_workbook([CHINA, MEXICO], tmp_path / "lc.xlsx")
    wb = load_workbook(out)
    ws = wb["China"]
    expected = compute_scenario(CHINA)["landed_unit"]
    # find the "= LANDED COST / unit" row in the build-up
    found = None
    for row in ws.iter_rows(min_col=1, max_col=2, values_only=True):
        if row[0] and "LANDED COST" in str(row[0]):
            found = row[1]
            break
    assert found is not None
    assert abs(float(found) - expected) < 1e-6


def test_comparison_flags_cheaper_total_despite_higher_ex_works(tmp_path):
    # Mexico ex-works ($46) > China ($42.50) but the 25% tariff should make
    # China's LANDED cost higher — the whole point of the workbook.
    out = build_landed_cost_workbook([CHINA, MEXICO], tmp_path / "lc.xlsx")
    wb = load_workbook(out)
    ws = wb["Comparison"]
    # scan all cells for the "Lowest landed cost" callout
    flat = [str(c) for row in ws.iter_rows(values_only=True) for c in row if c]
    assert any("Lowest landed cost: Mexico" in s for s in flat)


def test_sensitivity_reports_a_flip_point(tmp_path):
    out = build_landed_cost_workbook([CHINA, MEXICO], tmp_path / "lc.xlsx")
    wb = load_workbook(out)
    ws = wb["Sensitivity"]
    flat = [str(c) for row in ws.iter_rows(values_only=True) for c in row if c]
    # at 0% tariff China is cheaper (42.50 base < 46 base); at 25% it isn't →
    # somewhere in the sweep the winner flips
    assert any("flip" in s.lower() for s in flat)


def test_empty_scenarios_raise(tmp_path):
    try:
        build_landed_cost_workbook([], tmp_path / "x.xlsx")
        assert False, "expected ValueError"
    except ValueError:
        pass
