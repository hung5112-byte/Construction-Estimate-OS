"""End-to-end: the unattended pipeline on the synthetic bid package (no LLM anywhere)."""
from __future__ import annotations

import json
import shutil
from pathlib import Path

import pytest
from core.estimating import pipeline

REPO = Path(__file__).resolve().parents[3]
FIX = REPO / "docs" / "tests" / "fixtures" / "sample-set-prairie-creek"
GT = json.loads((FIX / "ground_truth.json").read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def run(tmp_path_factory):
    vault = tmp_path_factory.mktemp("vault")
    shutil.copytree(REPO / "00-Brain", vault / "00-Brain")
    folder = pipeline.run_all(vault, FIX, "Prairie Creek Office/Warehouse Bldg 2", "office-warehouse", "Plano", 2)
    return vault, folder


def _j(folder: Path, name: str):
    return json.loads((folder / name).read_text(encoding="utf-8"))


def test_folder_has_every_stage_file(run):
    _, folder = run
    for f in ("00-project-profile.md", "01-sheet-register.md", "02-spec-index.md", "03-sheets.json", "04-takeoff-ledger.md",
              "05-clarification.md", "06-estimate.md", "07-review-scorecard.md", "08-estimate-report.md"):
        assert (folder / f).exists(), f
    assert len(list((folder / "cards").glob("*.md"))) == len(GT["sheets"])
    assert len(list((folder / "package").glob("*.pdf"))) == len(GT["sheets"]) + 1


def test_profile_from_code_summary_and_spec(run):
    _, folder = run
    p = _j(folder, "00-project-profile.json")
    assert p["gross_sf"] == GT["gross_sf"] and p["areas"]["OFFICE AREA"] == GT["office_sf"] and p["heights_ft"] == [14.0, 24.0]
    assert p["bid_due"] == GT["spec"]["bid_due"] and p["rfi_cutoff"] == GT["spec"]["rfi_cutoff"] and p["duration_days"] == GT["spec"]["duration_days"]
    assert {"31", "32", "33"} <= set(p["exclusions"])           # site work by separate contract, cited


def test_seed_takeoff_matches_ground_truth_counts(run):
    _, folder = run
    items = _j(folder, "04-takeoff-ledger.json")
    by_code: dict[str, float] = {}
    for it in items:
        by_code[it["item_code"]] = by_code.get(it["item_code"], 0) + it["qty"]
    doors = sum(v for k, v in by_code.items() if k.startswith("08-hm-door") or k == "08-storefront-entrance-pair")
    assert doors == GT["door_count"] and by_code["08-overhead-door-12x14"] == len(GT["overhead_doors"])
    assert by_code["08-window-fixed-aluminum"] == GT["window_count"]
    assert by_code["03-footing-spread"] == pytest.approx(GT["footings"]["F1_count"] * GT["footings"]["F1_cy_each"], rel=1e-3)
    assert by_code["03-footing-continuous"] == pytest.approx(GT["footings"]["F2_cy"], abs=0.1)
    assert by_code["03-sog-4in"] == GT["office_sf"] and by_code["03-sog-6in"] == GT["warehouse_sf"]
    assert by_code["05-structural-steel"] + by_code["05-steel-joists"] == pytest.approx(GT["steel_total_tons"], abs=0.05)
    assert by_code["05-metal-deck"] == GT["deck_sf"]
    assert by_code["21-sprinkler-design-build"] == GT["sprinkler_heads"]
    assert by_code["22-water-closet"] == GT["plumbing_fixtures"]["WC-1"] and by_code["22-hose-bibb"] == GT["plumbing_fixtures"]["HB-1"]
    assert by_code["26-troffer-led-2x4"] == GT["electrical"]["fixtures"]["A"] and by_code["26-receptacle-duplex"] == GT["electrical"]["receptacles"]
    assert by_code["28-facp"] == 1 and by_code["28-fa-device"] == 20
    assert by_code["23-rtu-7-5-ton"] == 2 and by_code["23-diffuser-grille"] == GT["diffusers"] + 6
    assert by_code["09-act-ceiling-2x2"] == sum(r["area_sf"] for r in GT["rooms"] if r["number"] not in ("106", "107", "200"))
    assert by_code["10-signage-allowance"] == 1
    checks = _j(folder, "03-takeoff-checks.json")
    assert all(c["status"] == "match" for c in checks) and len(checks) == 3
    assert _j(folder, "04-discrepancies.json") == []
    assert all(it["sheet"] and it["method"] and 0 <= it["confidence"] <= 1 for it in items)


def test_pricing_gates_and_outputs(run):
    vault, folder = run
    est = _j(folder, "06-estimate.json")
    assert not [p for p in est["lines"] if "UNPRICED" in p["flags"]]
    assert est["meta"]["location_factor"] == 0.88
    pol = est["meta"]["policy"]
    assert pol["contingency_pct"] == 3.0 and "markup-policy.md" in pol["contingency_basis"] and pol["fee_pct"] == 5.0
    assert est["markups"][-1]["name"] == "TOTAL BID" and est["markups"][-1]["amount"] > 1_000_000
    bench = est["meta"]["benchmark"]
    assert 140 <= bench["per_sf"] <= 260 and bench["checks"][0]["status"] == "inside"
    sc = _j(folder, "07-review-scorecard.json")
    assert sc["verdict"] == "APPROVE", [g for g in sc["gates"] if not g["passed"]]
    assert "excluded with citation: 31" in next(g for g in sc["gates"] if g["id"] == "G3")["details"]
    qs = _j(folder, "05-questions.json")
    assert qs and all(q["answer"] or q["assumption"] for q in qs)      # auto-assumed, never silent
    report = (folder / "08-estimate-report.md").read_text(encoding="utf-8")
    assert "verdict: APPROVE" in report and "AUTO-ASSUMED" in report and "Excluded — division 31" in report
    outs = vault / "03-Outputs" / folder.name
    assert (outs / "estimate-workbook.xlsx").stat().st_size > 10_000 and (outs / "basis-of-estimate.docx").stat().st_size > 5_000
    from openpyxl import load_workbook
    wb = load_workbook(outs / "estimate-workbook.xlsx", read_only=True)
    assert {"Summary", "Priced lines", "Takeoff ledger", "General conditions", "Review gates", "RFI log", "Sheet register", "Sources"} <= set(wb.sheetnames)


def test_reader_merge_reconciles_against_schedule(run, tmp_path):
    _, folder = run
    reader = tmp_path / "reader.json"
    reader.write_text(json.dumps([
        {"division": "08", "item_code": "08-hm-door-single", "description": "HM door 102", "qty": 1, "unit": "EA", "sheet": "A-601", "method": "vision", "discipline": "03-architectural", "confidence": 0.7, "tags": {"mark": "102"}},
        {"division": "09", "item_code": "09-paint-cmu", "description": "Paint exposed CMU warehouse", "qty": 6240, "unit": "SF", "sheet": "A-201", "method": "vision", "discipline": "03-architectural", "confidence": 0.75},
    ]), encoding="utf-8")
    before = len(_j(folder, "04-takeoff-ledger.json"))
    led = pipeline.merge_reader_lines(folder, reader)
    codes = [i.item_code for i in led.items]
    assert "09-paint-cmu" in codes and len(led.items) == before + 1      # the door 102 vision line collapsed into the schedule line
    door102 = [i for i in led.items if i.tags.get("mark") == "102" and i.item_code == "08-hm-door-single"]
    assert len(door102) == 1 and door102[0].method == "schedule"
