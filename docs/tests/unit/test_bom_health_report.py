from openpyxl import load_workbook

from core.tools.bom_health import BomHealth
from core.tools.bom_health_report import build_bom_health_report, score_part

BOM = [
    {"part": "IC-9", "lifecycle": "eol", "sources": 1, "lead_time": 30, "coverage": 4, "ltb_days": 45},
    {"part": "R-1234", "lifecycle": "active", "sources": 3, "lead_time": 8, "coverage": 16, "compliance": "ok"},
    {"part": "U-77", "lifecycle": "obsolete", "sources": 1},
    {"part": "C-55", "lifecycle": "active", "sources": 1, "lead_time": 10, "coverage": 14},
]


def test_score_part_matches_tool():
    tool = BomHealth().run("bom IC-9 lifecycle=eol sources=1 lead_time=30 coverage=4 ltb_days=45").data
    rep = score_part(BOM[0])
    assert rep["risk_score"] == tool["risk_score"]
    assert rep["disposition"] == tool["disposition"]


def test_report_has_both_sheets(tmp_path):
    out = build_bom_health_report(BOM, tmp_path / "bom.xlsx")
    assert out.exists()
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "Parts"]


def test_parts_sheet_is_sorted_worst_first(tmp_path):
    out = build_bom_health_report(BOM, tmp_path / "bom.xlsx")
    wb = load_workbook(out)
    ws = wb["Parts"]
    # column 7 = Risk; rows start at 5
    risks = [ws.cell(r, 7).value for r in range(5, 5 + len(BOM))]
    assert risks == sorted(risks, reverse=True)


def test_summary_counts_single_source_and_eol(tmp_path):
    out = build_bom_health_report(BOM, tmp_path / "bom.xlsx")
    wb = load_workbook(out)
    ws = wb["Summary"]
    flat = {}
    for row in ws.iter_rows(min_row=5, max_col=2, values_only=True):
        if row[0]:
            flat[row[0]] = row[1]
    assert flat["Total parts"] == 4
    assert flat["Single-source parts"] == 3   # IC-9, U-77, C-55
    assert flat["EOL / obsolete parts"] == 2   # IC-9 (EOL), U-77 (obsolete)


def test_empty_bom_raises(tmp_path):
    try:
        build_bom_health_report([], tmp_path / "x.xlsx")
        assert False, "expected ValueError"
    except ValueError:
        pass
