"""Compliance matrix report generator.

Scores a list of certification/compliance items through the `compliance_tracker`
tool and writes a filed 2-sheet .xlsx:

  • Summary — counts by status, shipment-blocking count (required & not shippable),
              expiring-soon count, and the blocking items first
  • Items   — one row per product×cert (status, days-to-expiry, tier, shippable,
              action), worst first

Every value is computed by `compliance_tracker`, so the matrix can never disagree
with the tool.
"""
from __future__ import annotations
from datetime import date
from pathlib import Path

from core.tools.compliance_tracker import ComplianceTracker, _DISCLAIMER

_CT = ComplianceTracker()

_QUERY_KEYS = ("status", "expiry", "days_to_expiry", "required", "lead_time_days", "as_of")
_TIER_RANK = {"Critical": 0, "High": 1, "Moderate": 2, "Low": 3}
_ITEMS_HEADER = (
    "Product", "Cert", "Status", "Days to expiry", "Tier", "Shippable", "Action",
)


def score_item(item: dict) -> dict:
    return _CT.run(_item_query(item)).data


def build_compliance_matrix(items: list[dict], out_path: str | Path, *, as_of: str | None = None) -> Path:
    from openpyxl import Workbook

    if not items:
        raise ValueError("need at least one compliance item")
    as_of = as_of or date.today().isoformat()
    scored = [(it, score_item({**it, "as_of": it.get("as_of", as_of)})) for it in items]
    scored.sort(key=lambda it: (
        _TIER_RANK.get(it[1].get("risk_tier"), 9),
        it[1].get("days_to_expiry") if it[1].get("days_to_expiry") is not None else 10**9,
    ))

    wb = Workbook()
    wb.remove(wb.active)
    _sheet_summary(wb, scored, as_of)
    _sheet_items(wb, scored, as_of)

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))
    return out_path


def _sheet_summary(wb, scored, as_of):
    ws = wb.create_sheet("Summary")
    _title(ws, "Compliance Matrix Summary", as_of, ncols=2)
    statuses = [d.get("status") for _, d in scored]
    blocking = [(it, d) for it, d in scored if d.get("required") and not d.get("shippable")]
    rows = [
        ("Total items", len(scored)),
        ("Shipment-blocking", len(blocking)),
        ("Expiring soon", statuses.count("Expiring soon")),
        ("Expired", statuses.count("Expired")),
        ("Missing", statuses.count("Missing")),
        ("Non-compliant", statuses.count("Non-compliant")),
        ("Valid", statuses.count("Valid")),
    ]
    _header_row(ws, 4, ["Metric", "Value"])
    r = 5
    for label, val in rows:
        ws.cell(r, 1, label)
        cell = ws.cell(r, 2, val)
        if label == "Shipment-blocking" and val:
            cell.font = _bold()
        r += 1
    r += 1
    ws.cell(r, 1, "Shipment-blocking items").font = _bold()
    r += 1
    _header_row(ws, r, ["Product", "Cert", "Status", "Action"])
    r += 1
    for it, d in blocking[:10]:
        ws.cell(r, 1, str(it.get("product", "")))
        ws.cell(r, 2, str(d.get("cert", it.get("cert", "?"))))
        ws.cell(r, 3, d.get("status", ""))
        ws.cell(r, 4, d.get("recommended_action", ""))
        r += 1
    if not blocking:
        ws.cell(r, 1, "None — nothing is blocking shipment.").font = _italic()
        r += 1
    _note(ws, r + 1, _DISCLAIMER, ncols=4)
    _autofit(ws, {1: 18, 2: 12, 3: 14, 4: 60})


def _sheet_items(wb, scored, as_of):
    ws = wb.create_sheet("Items")
    _title(ws, "Compliance Items — worst first", as_of, ncols=len(_ITEMS_HEADER))
    _header_row(ws, 4, list(_ITEMS_HEADER))
    r = 5
    for it, d in scored:
        ws.cell(r, 1, str(it.get("product", "")))
        ws.cell(r, 2, str(d.get("cert", it.get("cert", "?"))))
        ws.cell(r, 3, d.get("status", ""))
        dte = d.get("days_to_expiry")
        ws.cell(r, 4, "" if dte is None else dte)
        ws.cell(r, 5, d.get("risk_tier", ""))
        ws.cell(r, 6, "yes" if d.get("shippable") else "NO")
        ws.cell(r, 7, d.get("recommended_action", ""))
        r += 1
    _note(ws, r + 1, _DISCLAIMER, ncols=len(_ITEMS_HEADER))
    _autofit(ws, {1: 18, 2: 12, 3: 14, 7: 52})
    ws.freeze_panes = "C5"


# ---- helpers ----
def _item_query(item: dict) -> str:
    cert = item.get("cert", item.get("name", ""))
    parts = [f"cert {cert}".strip()]
    for k in _QUERY_KEYS:
        if item.get(k) is not None:
            parts.append(f"{k}={item[k]}")
    return " ".join(str(p) for p in parts)


def _bold():
    from openpyxl.styles import Font
    return Font(bold=True)


def _italic():
    from openpyxl.styles import Font
    return Font(italic=True, color="555555")


def _title(ws, text, as_of, ncols):
    from openpyxl.styles import Font
    ws.cell(1, 1, text).font = Font(bold=True, size=14)
    ws.cell(2, 1, f"As of {as_of} · general info — verify certs of record with compliance").font = _italic()
    if ncols > 1:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)


def _header_row(ws, row, headers):
    from openpyxl.styles import Font, PatternFill, Alignment
    fill = PatternFill("solid", fgColor="1F4E78")
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="left" if c == 1 else "center")


def _note(ws, row, text, ncols):
    cell = ws.cell(row, 1, f"⚠ {text}")
    cell.font = _italic()
    if ncols > 1:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)


def _autofit(ws, widths):
    from openpyxl.utils import get_column_letter
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
