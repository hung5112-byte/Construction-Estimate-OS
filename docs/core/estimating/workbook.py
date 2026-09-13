"""Estimate workbook (.xlsx) — every number computed by the cost engine, rendered with openpyxl."""
from __future__ import annotations

from pathlib import Path

_MONEY = '$#,##0.00'
_PCT = '0.0"%"'


def _autosize(ws, widths: dict[int, int] | None = None) -> None:
    from openpyxl.utils import get_column_letter

    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        best = 10
        for cell in ws[letter]:
            if cell.value is not None:
                best = max(best, min(60, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = (widths or {}).get(col, best)


def _header(ws, row: int, values: list[str]) -> None:
    from openpyxl.styles import Font, PatternFill

    for c, v in enumerate(values, 1):
        cell = ws.cell(row, c, v)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")


def write_workbook(out_path: Path, estimate: dict, ledger_items: list[dict], gates: list[dict], questions: list[dict],
                   register: list[dict], profile: dict, risks: list[dict] | None = None) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    # ---- Summary
    ws = wb.active
    ws.title = "Summary"
    from openpyxl.styles import Font as _Font

    ws["A1"] = f"ESTIMATE SUMMARY — {profile.get('name', '')}"
    ws["A1"].font = _Font(bold=True, size=14)
    ws["A2"] = f"Building type: {profile.get('building_type', '')} · GSF {profile.get('gross_sf', 0):,.0f} · AACE Class {profile.get('aace_class', '')} · {profile.get('city', '')}"
    ws["A3"] = estimate["meta"].get("disclaimer", "")
    _header(ws, 5, ["Division", "Title", "Labor", "Material", "Equipment", "Subcontract", "Total", "$/SF", "% of direct", "Lines", "Unpriced"])
    r = 6
    for row in estimate["summary_by_division"]:
        vals = [row["division"], row["title"], row["labor"], row["material"], row["equipment"], row["sub"], row["total"], row["per_sf"], row["pct"], row["lines"], row["unpriced"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            if 3 <= c <= 8:
                cell.number_format = _MONEY
            if c == 9:
                cell.number_format = _PCT
        r += 1
    r += 1
    _header(ws, r, ["Markup line", "%", "Base", "Amount", "Basis"])
    r += 1
    for m in estimate["markups"]:
        ws.cell(r, 1, m["name"])
        ws.cell(r, 2, m["pct"]).number_format = _PCT
        ws.cell(r, 3, m["base"]).number_format = _MONEY
        ws.cell(r, 4, m["amount"]).number_format = _MONEY
        ws.cell(r, 5, m["basis"])
        r += 1
    total = next((m["amount"] for m in estimate["markups"] if m["name"] == "TOTAL BID"), None)
    if total and profile.get("gross_sf"):
        r += 1
        ws.cell(r, 1, "TOTAL BID $/SF")
        ws.cell(r, 4, round(total / profile["gross_sf"], 2)).number_format = _MONEY
    _autosize(ws)
    # ---- Lines
    ws = wb.create_sheet("Priced lines")
    _header(ws, 1, ["ID", "Div", "Item code", "Description", "Qty", "Unit", "Unit L", "Unit M", "Unit E", "Unit Sub", "Labor", "Material", "Equipment", "Sub", "Total", "Source", "Sheet", "Method", "Conf", "Flags"])
    for i, p in enumerate(estimate["lines"], 2):
        vals = [p["item_id"], p["division"], p["item_code"], p["description"], p["qty"], p["unit"], p["unit_labor"], p["unit_material"], p["unit_equipment"], p["unit_sub"],
                p["labor"], p["material"], p["equipment"], p["sub"], p["total"], p["source"], p["sheet"], p["method"], p["confidence"], ", ".join(p["flags"])]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(i, c, v)
            if 7 <= c <= 15:
                cell.number_format = _MONEY
    _autosize(ws)
    # ---- Takeoff
    ws = wb.create_sheet("Takeoff ledger")
    _header(ws, 1, ["ID", "Div", "Item code", "Description", "Qty", "Unit", "Waste %", "Sheet", "Rev", "Method", "Confidence", "Notes"])
    for i, it in enumerate(ledger_items, 2):
        for c, v in enumerate([it["id"], it["division"], it["item_code"], it["description"], it["qty"], it["unit"], it["waste_pct"], it["sheet"], it.get("revision") or "", it["method"], it["confidence"], it.get("notes", "")], 1):
            ws.cell(i, c, v)
    _autosize(ws)
    # ---- General conditions
    ws = wb.create_sheet("General conditions")
    _header(ws, 1, ["Line", "Qty", "Unit", "Rate", "Total", "Basis"])
    for i, g in enumerate(estimate["general_conditions"], 2):
        ws.cell(i, 1, g["name"])
        ws.cell(i, 2, g["qty"])
        ws.cell(i, 3, g["unit"])
        ws.cell(i, 4, g["rate"]).number_format = _MONEY
        ws.cell(i, 5, g["total"]).number_format = _MONEY
        ws.cell(i, 6, g["basis"])
    _autosize(ws)
    # ---- Review gates
    ws = wb.create_sheet("Review gates")
    _header(ws, 1, ["Gate", "Check", "Passed", "Blocking", "Details", "Evidence"])
    for i, g in enumerate(gates, 2):
        for c, v in enumerate([g["id"], g["name"], "PASS" if g["passed"] else "FAIL", "yes" if g["blocking"] else "no", g["details"], "; ".join(g.get("evidence", []))[:500]], 1):
            ws.cell(i, c, v)
    _autosize(ws)
    # ---- RFI log
    ws = wb.create_sheet("RFI log")
    _header(ws, 1, ["#", "Severity", "Question", "Citation", "Cost exposure", "Answer / assumption", "Status"])
    for i, q in enumerate(questions, 2):
        for c, v in enumerate([i - 1, q.get("severity", ""), q.get("text", ""), q.get("citation", ""), q.get("exposure", ""), q.get("answer") or q.get("assumption") or "", "answered" if q.get("answer") else ("assumed" if q.get("assumption") else "open")], 1):
            ws.cell(i, c, v)
    _autosize(ws)
    # ---- Sheet register
    ws = wb.create_sheet("Sheet register")
    _header(ws, 1, ["Sheet", "Title", "Discipline", "Type", "Scale", "Rev", "Date", "Size", "Vector", "Reader dept"])
    for i, s in enumerate(register, 2):
        for c, v in enumerate([s["sheet_id"], s["title"], s["discipline_name"], s["sheet_type"], s.get("scale_label") or "", s.get("revision") or "", s.get("date") or "", s["size_name"], "yes" if s["is_vector"] else "scanned", s["reader_department"]], 1):
            ws.cell(i, c, v)
    _autosize(ws)
    # ---- Risks
    ws = wb.create_sheet("Risk register")
    _header(ws, 1, ["Risk", "Probability", "Impact $", "EMV $", "Owner", "Response"])
    for i, rk in enumerate(risks or [], 2):
        for c, v in enumerate([rk.get("risk"), rk.get("probability"), rk.get("impact"), rk.get("emv"), rk.get("owner"), rk.get("response")], 1):
            ws.cell(i, c, v)
    _autosize(ws)
    # ---- Sources
    ws = wb.create_sheet("Sources")
    _header(ws, 1, ["Source", "Used by (lines)"])
    counts: dict[str, int] = {}
    for p in estimate["lines"]:
        counts[p["source"] or "(unpriced)"] = counts.get(p["source"] or "(unpriced)", 0) + 1
    for i, (src, n) in enumerate(sorted(counts.items(), key=lambda kv: -kv[1]), 2):
        ws.cell(i, 1, src)
        ws.cell(i, 2, n)
    ws.cell(i + 2, 1, estimate["meta"].get("disclaimer", ""))
    _autosize(ws)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))
    return out_path
