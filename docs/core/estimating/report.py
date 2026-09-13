"""Estimate report skeleton — verdict, numbers, blockers and open questions are assembled in code;
the writer agent (or nobody, in --no-llm mode) adds narrative under the marked sections."""
from __future__ import annotations

from pathlib import Path


def report_markdown(profile: dict, estimate: dict, gates: list[dict], questions: list[dict], benchmark: dict,
                    ledger_items: list[dict], narrative: dict | None = None) -> str:
    narrative = narrative or {}
    markups = {m["name"]: m for m in estimate["markups"]}
    total = markups.get("TOTAL BID", {}).get("amount", 0.0)
    direct = markups.get("Direct cost", {}).get("amount", 0.0)
    gsf = float(profile.get("gross_sf") or 0)
    per_sf = total / gsf if gsf else None
    v = "REVISE" if any(g["blocking"] and not g["passed"] for g in gates) else "APPROVE"
    cls = profile.get("aace_class", 2)
    band = {5: "−30% / +50%", 4: "−20% / +30%", 3: "−15% / +20%", 2: "−10% / +15%", 1: "−5% / +10%"}.get(int(cls), "")
    open_q = [q for q in questions if not q.get("answer") and not q.get("assumption")]
    crit = [q for q in open_q if q.get("severity") == "CRITICAL"]
    low = [it for it in ledger_items if it["confidence"] < 0.6]
    unpriced = [p for p in estimate["lines"] if "UNPRICED" in p.get("flags", [])]
    uncertain = sum(1 for p in estimate["lines"] if "UNCERTAIN" in p.get("flags", []))
    out = ["---", "type: estimate_report", "stop: 1", f"verdict: {v}", f"aace_class: {cls}", "---",
           f"# Estimate report — {profile.get('name', '')}", "",
           "## TL;DR", narrative.get("tldr") or (
               f"{profile.get('building_type', 'building')} of {gsf:,.0f} SF in {profile.get('city', '')}: **${total:,.0f}** total bid "
               f"(${per_sf:,.2f}/SF)" + (", inside the benchmark band" if benchmark.get('checks') and benchmark['checks'][0]['status'] == 'inside' else "") +
               f". Class {cls} estimate, expected accuracy {band}. Review verdict **{v}**; {len(crit)} open CRITICAL question(s); "
               f"{len(unpriced)} unpriced line(s); {uncertain} line(s) priced from placeholder data."), "",
           "## Verdict", f"**{v}** — " + (narrative.get("verdict") or ("all blocking gates passed" if v == "APPROVE" else "one or more blocking gates failed — see scorecard")), "",
           "## Estimate summary", "", "| Division | Title | Total | $/SF | % of direct |", "|---|---|---|---|---|"]
    for r in estimate["summary_by_division"]:
        out.append(f"| {r['division']} | {r['title']} | ${r['total']:,.0f} | {r['per_sf'] if r['per_sf'] is not None else ''} | {r['pct']}% |")
    out += ["", "| Markup | % | Amount |", "|---|---|---|"]
    for m in estimate["markups"]:
        out.append(f"| {m['name']} | {m['pct'] if m['pct'] else ''} | ${m['amount']:,.0f} |")
    out += ["", f"**Direct cost ${direct:,.0f} · Total bid ${total:,.0f} · ${per_sf:,.2f}/SF**" if per_sf else f"**Total bid ${total:,.0f}**", ""]
    out += ["## Benchmark position", ""]
    for c in benchmark.get("checks", []):
        out.append(f"- {c['check']}: **{c['value']}** vs {c['band']} → {c['status']} _(source: {c['source']})_")
    out += ["", "## What was measured vs assumed", ""]
    by_method: dict[str, int] = {}
    for it in ledger_items:
        by_method[it["method"]] = by_method.get(it["method"], 0) + 1
    out.append("- Lines by method: " + ", ".join(f"{k} {n}" for k, n in sorted(by_method.items())))
    out.append(f"- Lines below the confidence floor (verify before bid): {len(low)}")
    for it in low[:15]:
        out.append(f"  - {it['id']} {it['description']} — {it['qty']:,.1f} {it['unit']} ({it['sheet']}, {it['method']}, conf {it['confidence']:.2f}) {it.get('notes','')}")
    out.append(narrative.get("measured_vs_assumed", ""))
    out += ["", "## Open questions and cost exposure", ""]
    if not open_q:
        out.append("- none open")
    for q in open_q[:25]:
        out.append(f"- [{q.get('severity','')}] {q.get('text','')} _(cite: {q.get('citation','')}; exposure: {q.get('exposure','n/a')})_")
    out += ["", "## Blockers", ""]
    blockers = [g for g in gates if g["blocking"] and not g["passed"]]
    out.extend([f"- {g['id']} {g['name']}: {g['details']}" for g in blockers] or ["- none"])
    out += ["", "## Assumptions and exclusions", ""]
    for a in (narrative.get("assumptions") or profile.get("assumptions") or ["Standing exclusions per 00-Brain/markup-policy.md apply unless the bid form says otherwise."]):
        out.append(f"- {a}")
    out += ["", "## Review scorecard", "", "| Gate | Result | Details |", "|---|---|---|"]
    for g in gates:
        out.append(f"| {g['id']} {g['name']} | {'pass' if g['passed'] else ('FAIL' if g['blocking'] else 'warn')} | {g['details']} |")
    out += ["", "## Sources", "", f"- Cost library: {estimate['meta'].get('library_paths', '')}",
            f"- Location factor: {estimate['meta'].get('location_factor', '')} ({estimate['meta'].get('location_source', '')})",
            f"- {estimate['meta'].get('disclaimer', '')}", "",
            "## Narrative", narrative.get("narrative") or "_(no writer agent ran — deterministic skeleton only)_", ""]
    return "\n".join(out)


def write_basis_of_estimate_docx(out_path: Path, report_md: str, title: str) -> Path:
    """Render the report markdown into a simple .docx (python-docx) — headings, bullets, tables as text."""
    from docx import Document

    doc = Document()
    doc.add_heading(title, 0)
    table_buf: list[str] = []

    def flush_table():
        nonlocal table_buf
        rows = [r for r in table_buf if not set(r.replace("|", "").strip()) <= {"-", " "}]
        if rows:
            cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
            ncols = max(len(r) for r in cells)
            t = doc.add_table(rows=0, cols=ncols)
            t.style = "Table Grid"
            for r in cells:
                cells_row = t.add_row().cells
                for i, c in enumerate(r):
                    cells_row[i].text = c
        table_buf = []

    in_fm = False
    for line in report_md.splitlines():
        if line.strip() == "---":
            in_fm = not in_fm
            continue
        if in_fm:
            continue
        if line.startswith("|"):
            table_buf.append(line)
            continue
        if table_buf:
            flush_table()
        if line.startswith("# "):
            doc.add_heading(line[2:], 1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], 2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:].replace("**", ""), style="List Bullet")
        elif line.strip():
            doc.add_paragraph(line.replace("**", "").replace("_", ""))
    if table_buf:
        flush_table()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
